from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

wb = load_workbook("./RPAbasic/excel/range.xlsx")
ws = wb.active

# 막대차트 ( 차트 범위 설정 )
bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)

# 차트 종류 설정
bar_chart = BarChart()
bar_chart.add_data(bar_value)

# 차트 삽입
ws.add_chart(bar_chart, "E1")

wb.save("./RPAbasic/excel/range_chart.xlsx")
